import io
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from api.rag import DocumentError, RagStore, _chunk_text


def text_pdf(text):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class RagStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.store = RagStore(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pdf_ingestion_keeps_page_citation_and_evidence(self):
        document = self.store.add_document(
            "owner-a", "health-policy.pdf",
            text_pdf("The annual deductible is 1200 dollars for each covered member."),
        )
        result = self.store.answer("owner-a", "What is the annual deductible?")
        self.assertEqual(document["status"], "ready")
        self.assertEqual(result["retrieved_count"], 1)
        self.assertEqual(result["sources"][0]["filename"], "health-policy.pdf")
        self.assertEqual(result["sources"][0]["page"], 1)
        self.assertIn("1200 dollars", result["answer"])
        self.assertIn("[1]", result["answer"])
        self.assertIn("1200 dollars", result["sources"][0]["excerpt"])

    def test_index_and_embeddings_survive_store_reopen(self):
        self.store.add_document(
            "owner-a", "persisted.txt",
            b"The annual out-of-pocket maximum is 5000 dollars.",
        )
        reopened = RagStore(self.temp_dir.name)
        result = reopened.answer("owner-a", "What is the out-of-pocket maximum?")
        self.assertIn("5000 dollars", result["answer"])
        self.assertEqual(result["sources"][0]["filename"], "persisted.txt")

    def test_txt_sections_and_multidocument_filtering(self):
        first = self.store.add_document(
            "owner-a", "terms.txt",
            b"Section 4 Waiting period\nThe waiting period for maternity care is 24 months.",
        )
        self.store.add_document(
            "owner-a", "claims.txt",
            b"Section 8 Claims\nSubmit a claim with an itemized bill within 30 days.",
        )
        result = self.store.answer(
            "owner-a", "What is the maternity waiting period?", [first["id"]]
        )
        self.assertEqual(result["sources"][0]["document_id"], first["id"])
        self.assertEqual(result["sources"][0]["section"], "Section 4 Waiting period")
        self.assertIn("24 months", result["answer"])

    def test_docx_heading_and_retrieval_evaluation_set(self):
        document = Document()
        document.add_heading("Benefits", level=1)
        document.add_paragraph("Routine eye examinations are covered once every calendar year.")
        buffer = io.BytesIO()
        document.save(buffer)
        coverage = self.store.add_document("owner-a", "benefits.docx", buffer.getvalue())
        claims = self.store.add_document(
            "owner-a", "claims.txt",
            b"Claims procedure\nSubmit claims with original receipts within 45 days of treatment.",
        )
        exclusions = self.store.add_document(
            "owner-a", "exclusions.txt",
            b"Exclusions\nCosmetic surgery is excluded unless medically necessary after an accident.",
        )

        evaluation_cases = [
            ("How often are eye exams covered?", coverage["id"], "calendar year"),
            ("What is the deadline to submit a claim?", claims["id"], "45 days"),
            ("Is cosmetic surgery excluded?", exclusions["id"], "excluded"),
        ]
        for question, expected_document_id, expected_evidence in evaluation_cases:
            with self.subTest(question=question):
                result = self.store.answer("owner-a", question)
                self.assertTrue(result["sources"])
                self.assertEqual(result["sources"][0]["document_id"], expected_document_id)
                self.assertIn(expected_evidence, result["sources"][0]["excerpt"])
        coverage_result = self.store.answer("owner-a", evaluation_cases[0][0])
        self.assertEqual(coverage_result["sources"][0]["section"], "Benefits")

        combined = self.store.answer(
            "owner-a",
            "What coverage and claim submission deadlines do these documents describe "
            "for eye examinations and original receipts?",
        )
        self.assertEqual(
            {source["document_id"] for source in combined["sources"]},
            {coverage["id"], claims["id"]},
        )
        unsupported = self.store.answer("owner-a", "Does this cover lunar teleportation?")
        self.assertEqual(unsupported["sources"], [])

    def test_docx_tables_are_extracted(self):
        document = Document()
        document.add_heading("Limits", level=1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Emergency ambulance"
        table.cell(0, 1).text = "Up to 5000 dollars per year"
        buffer = io.BytesIO()
        document.save(buffer)
        self.store.add_document("owner-a", "limits.docx", buffer.getvalue())
        result = self.store.answer("owner-a", "What is the ambulance coverage limit?")
        self.assertTrue(result["sources"])
        self.assertEqual(result["sources"][0]["section"], "Limits")
        self.assertIn("5000 dollars", result["sources"][0]["excerpt"])

    def test_no_evidence_and_document_owner_isolation(self):
        document = self.store.add_document(
            "owner-a", "policy.txt", b"Routine dental care has a 60 dollar co-payment."
        )
        result = self.store.answer("owner-a", "What is the orbital launch protocol?")
        self.assertEqual(result["sources"], [])
        self.assertIn("do not provide sufficient information", result["answer"])
        self.assertEqual(self.store.list_documents("owner-b"), [])
        with self.assertRaises(DocumentError):
            self.store.answer("owner-b", "dental co-payment", [document["id"]])
        self.assertFalse(self.store.delete_document("owner-b", document["id"]))

    def test_delete_removes_document_and_chunks(self):
        document = self.store.add_document(
            "owner-a", "policy.txt", b"The policy provides emergency ambulance coverage."
        )
        stored_file = next(self.store.upload_dir.glob(f"{document['id']}.*"))
        self.assertTrue(stored_file.exists())
        self.assertTrue(self.store.delete_document("owner-a", document["id"]))
        self.assertFalse(stored_file.exists())
        self.assertEqual(self.store.list_documents("owner-a"), [])
        self.assertEqual(self.store.answer("owner-a", "ambulance coverage")["sources"], [])

    def test_rejects_unsupported_and_oversized_files(self):
        with self.assertRaises(DocumentError):
            self.store.add_document("owner-a", "script.exe", b"not a policy")
        with self.assertRaises(DocumentError):
            self.store.add_document(
                "owner-a", "large.txt", b"x" * (20 * 1024 * 1024 + 1)
            )

    def test_invalid_pdf_reports_failed_processing_status(self):
        with self.assertRaises(DocumentError) as error:
            self.store.add_document("owner-a", "broken.pdf", b"not a PDF")
        document = self.store.document_status("owner-a", error.exception.document_id)
        self.assertEqual(document["status"], "failed")
        self.assertIn("valid", document["error"])

    def test_chunking_preserves_overlap(self):
        chunks = _chunk_text("alpha " * 500)
        self.assertGreater(len(chunks), 1)
        self.assertLessEqual(len(chunks[0]["content"]), 1200)
        self.assertTrue(set(chunks[0]["content"].split()[-10:]) &
                        set(chunks[1]["content"].split()[:10]))


class RagApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from api import app as app_module

        cls.app_module = app_module
        cls.previous_store = app_module.RAG_STORE

    @classmethod
    def tearDownClass(cls):
        cls.app_module.RAG_STORE = cls.previous_store

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.app_module.RAG_STORE = RagStore(self.temp_dir.name)
        self.client = self.app_module.app.test_client()
        with self.client.session_transaction() as sess:
            sess["user"] = "admin"          # routes now require a logged-in session

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_upload_chat_and_isolated_document_library(self):
        upload = self.client.post(
            "/api/documents",
            data={"file": (io.BytesIO(
                b"Section 2 Coverage\nThe policy covers urgent ambulance transport."
            ), "policy.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(upload.status_code, 201)
        document = upload.get_json()["document"]
        response = self.client.post("/api/chat", json={
            "question": "What ambulance transport does the policy cover?",
            "document_ids": [document["id"]],
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("urgent ambulance transport", response.get_json()["answer"])
        self.assertEqual(response.get_json()["sources"][0]["section"], "Section 2 Coverage")
        self.assertEqual(
            self.client.get(f"/api/documents/{document['id']}/status").status_code, 200
        )

        other_client = self.app_module.app.test_client()
        with other_client.session_transaction() as sess:
            sess["user"] = "admin"
        self.assertEqual(other_client.get("/api/documents").get_json()["documents"], [])
        self.assertEqual(other_client.delete(f"/api/documents/{document['id']}").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/documents/{document['id']}").status_code, 200)

    def test_invalid_upload_and_bad_question_are_reported(self):
        invalid = self.client.post(
            "/api/documents",
            data={"file": (io.BytesIO(b"no"), "script.exe")},
            content_type="multipart/form-data",
        )
        self.assertEqual(invalid.status_code, 422)
        self.assertIn("PDF, TXT, or DOCX", invalid.get_json()["error"])
        response = self.client.post("/api/chat", json={"question": " "})
        self.assertEqual(response.status_code, 400)
        invalid_retrieval = self.client.post("/api/chat", json={
            "question": "policy", "top_k": {},
        })
        self.assertEqual(invalid_retrieval.status_code, 400)


if __name__ == "__main__":
    unittest.main()