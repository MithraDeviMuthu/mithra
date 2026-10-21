"""Local document retrieval for the insurance policy assistant."""
import json
import os
import re
import sqlite3
import uuid
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from sklearn.feature_extraction.text import HashingVectorizer
from werkzeug.utils import secure_filename


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_EXTRACTED_CHARS = 10_000_000
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200
DEFAULT_TOP_K = 5
MAX_TOP_K = 10
DEFAULT_MIN_SCORE = 0.08
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}
_vectorizer = HashingVectorizer(
    n_features=2**15,
    alternate_sign=False,
    norm="l2",
    ngram_range=(1, 2),
    stop_words="english",
    dtype=np.float32,
)


class DocumentError(ValueError):
    """An upload cannot be safely processed as a supported document."""


def _now():
    return datetime.now(timezone.utc).isoformat()


def _chunk_text(text, page=None, section=None):
    text = re.sub(r"[ \t]+", " ", text).strip()
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + CHUNK_SIZE // 2, end),
                           text.rfind(" ", start + CHUNK_SIZE // 2, end))
            if boundary > start:
                end = boundary
        content = text[start:end].strip()
        if content:
            chunks.append({"page": page, "section": section, "content": content})
        if end >= len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def _text_sections(text):
    sections = []
    heading = None
    lines = []
    heading_pattern = re.compile(
        r"^(?:#{1,6}\s+.+|(?:section|article|chapter)\s+[\w.-]+.*|\d+(?:\.\d+)*\s+\S.*)$",
        re.IGNORECASE,
    )
    for line in text.splitlines():
        stripped = line.strip()
        is_uppercase_heading = (
            len(stripped) <= 100 and any(char.isalpha() for char in stripped)
            and stripped == stripped.upper()
        )
        if stripped and (heading_pattern.match(stripped) or is_uppercase_heading):
            body = "\n".join(lines).strip()
            if body:
                sections.append((heading, body))
            heading = re.sub(r"^#{1,6}\s*", "", stripped)
            lines = []
        else:
            lines.append(line)
    body = "\n".join(lines).strip()
    if body:
        sections.append((heading, body))
    return sections or [(None, text)]


def _extract(path, extension):
    if extension == ".txt":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentError("TXT documents must be UTF-8 encoded.") from exc
        if len(text) > MAX_EXTRACTED_CHARS:
            raise DocumentError("Extracted document text exceeds the 10 million character limit.")
        return [
            chunk
            for section, body in _text_sections(text)
            for chunk in _chunk_text(body, section=section)
        ]

    if extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentError("PDF support requires the pypdf package.") from exc
        try:
            reader = PdfReader(str(path), strict=True)
            if len(reader.pages) > 1000:
                raise DocumentError("PDF documents are limited to 1,000 pages.")
            chunks = []
            extracted_chars = 0
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                extracted_chars += len(text)
                if extracted_chars > MAX_EXTRACTED_CHARS:
                    raise DocumentError("Extracted document text exceeds the 10 million character limit.")
                chunks.extend(_chunk_text(text, page=page_number))
            return chunks
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("The PDF could not be read. Check that it is a valid, unencrypted PDF.") from exc

    if extension == ".docx":
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocumentError("DOCX document expands beyond the 50 MB processing limit.")
                if "word/document.xml" not in archive.namelist():
                    raise DocumentError("The DOCX document is invalid.")
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except DocumentError:
            raise
        except (zipfile.BadZipFile, KeyError) as exc:
            raise DocumentError("The DOCX document is invalid.") from exc
        except ImportError as exc:
            raise DocumentError("DOCX support requires the python-docx package.") from exc
        try:
            document = Document(str(path))
            sections = []
            heading = None
            paragraphs = []
            extracted_chars = 0
            for block in document.element.body.iterchildren():
                if block.tag.endswith("}p"):
                    paragraph = Paragraph(block, document)
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    extracted_chars += len(text)
                    if extracted_chars > MAX_EXTRACTED_CHARS:
                        raise DocumentError("Extracted document text exceeds the 10 million character limit.")
                    if paragraph.style and paragraph.style.name.lower().startswith(("heading", "title")):
                        if paragraphs:
                            sections.append((heading, "\n".join(paragraphs)))
                        heading = text
                        paragraphs = []
                    else:
                        paragraphs.append(text)
                elif block.tag.endswith("}tbl"):
                    table = Table(block, document)
                    for row in table.rows:
                        text = " | ".join(
                            cell.text.strip().replace("\n", " / ")
                            for cell in row.cells if cell.text.strip()
                        )
                        if not text:
                            continue
                        extracted_chars += len(text)
                        if extracted_chars > MAX_EXTRACTED_CHARS:
                            raise DocumentError("Extracted document text exceeds the 10 million character limit.")
                        paragraphs.append(text)
            if paragraphs:
                sections.append((heading, "\n".join(paragraphs)))
            return [
                chunk
                for section, body in (sections or [(None, "")])
                for chunk in _chunk_text(body, section=section)
            ]
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("The DOCX document could not be read.") from exc

    raise DocumentError("Only PDF, TXT, and DOCX documents are supported.")


def _sparse_embedding(text):
    vector = _vectorizer.transform([text]).tocsr()
    return json.dumps(
        [[int(index), float(value)] for index, value in zip(vector.indices, vector.data)],
        separators=(",", ":"),
    )


def _terms(text):
    return set(re.findall(r"[a-z0-9]{2,}", text.lower())) - {
        "about", "after", "also", "does", "from", "have", "into", "must",
        "that", "their", "there", "these", "this", "what", "when", "where",
        "which", "with", "would",
    }


def _answer_from_evidence(question, sources):
    question_terms = _terms(question)
    excerpts = []
    for source in sources:
        sentences = re.split(r"(?<=[.!?])\s+|\n+", source["excerpt"])
        scored = [
            (len(question_terms & _terms(sentence)), len(sentence), sentence.strip())
            for sentence in sentences if sentence.strip()
        ]
        if scored:
            best = max(scored, key=lambda item: (item[0], -item[1]))
            excerpt = best[2] if best[0] else source["excerpt"].strip()
        else:
            excerpt = source["excerpt"].strip()
        if excerpt:
            excerpts.append(excerpt)
    if not excerpts:
        return "The available documents do not provide sufficient information to answer this question."
    return "The uploaded documents state:\n\n" + "\n\n".join(
        f'“{excerpt}” [{index}]' for index, excerpt in enumerate(excerpts, start=1)
    )


def _summarize(text, n=5):
    """Pick the n most representative sentences (extractive summary)."""
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+|\n+", text) if len(s.strip()) > 25]
    unique, seen = [], set()
    for s in sentences:                      # chunks overlap, so drop repeated sentences
        if s.lower() not in seen and not s.lower().startswith("summarize"):
            seen.add(s.lower())
            unique.append(s)
    freq = Counter(re.findall(r"[a-z0-9]{4,}", text.lower()))

    def score(s):
        words = re.findall(r"[a-z0-9]{4,}", s.lower())
        return sum(freq[w] for w in words) / (len(words) or 1)

    top = set(sorted(unique, key=score, reverse=True)[:n])
    return [s for s in unique if s in top]   # keep original order


class RagStore:
    def __init__(self, data_dir):
        self.data_dir = Path(data_dir).resolve()
        self.upload_dir = self.data_dir / "documents"
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.database = self.data_dir / "rag.sqlite3"
        self._initialize()

    @contextmanager
    def _connect(self):
        connection = sqlite3.connect(self.database)
        connection.row_factory = sqlite3.Row
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def _initialize(self):
        with self._connect() as connection:
            connection.executescript("""
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    owner_id TEXT NOT NULL,
                    filename TEXT NOT NULL,
                    stored_name TEXT NOT NULL,
                    status TEXT NOT NULL,
                    error TEXT,
                    page_count INTEGER,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS documents_owner_idx
                    ON documents(owner_id, created_at);
                CREATE TABLE IF NOT EXISTS chunks (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    page INTEGER,
                    section TEXT,
                    content TEXT NOT NULL,
                    embedding TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS chunks_document_idx ON chunks(document_id);
            """)
            connection.execute("PRAGMA foreign_keys=ON")

    def list_documents(self, owner_id):
        with self._connect() as connection:
            rows = connection.execute(
                """SELECT id, filename, status, error, page_count, chunk_count, created_at
                   FROM documents WHERE owner_id=? ORDER BY created_at DESC""",
                (owner_id,),
            ).fetchall()
        return [dict(row) for row in rows]

    def document_status(self, owner_id, document_id):
        with self._connect() as connection:
            row = connection.execute(
                """SELECT id, filename, status, error, page_count, chunk_count, created_at
                   FROM documents WHERE id=? AND owner_id=?""",
                (document_id, owner_id),
            ).fetchone()
        return dict(row) if row else None

    def add_document(self, owner_id, original_name, content):
        filename = secure_filename(Path(original_name or "").name)
        extension = Path(filename).suffix.lower()
        if not filename or extension not in ALLOWED_EXTENSIONS:
            raise DocumentError("Upload a PDF, TXT, or DOCX document.")
        if not content:
            raise DocumentError("The selected document is empty.")
        if len(content) > MAX_FILE_BYTES:
            raise DocumentError("Documents must be 20 MB or smaller.")

        document_id = uuid.uuid4().hex
        stored_name = f"{document_id}{extension}"
        path = self.upload_dir / stored_name
        with self._connect() as connection:
            connection.execute(
                """INSERT INTO documents
                   (id, owner_id, filename, stored_name, status, created_at)
                   VALUES (?, ?, ?, ?, 'processing', ?)""",
                (document_id, owner_id, filename, stored_name, _now()),
            )
        try:
            path.write_bytes(content)
            chunks = _extract(path, extension)
            if not chunks:
                raise DocumentError("No readable text was found in the document.")
            with self._connect() as connection:
                connection.execute("PRAGMA foreign_keys=ON")
                connection.executemany(
                    """INSERT INTO chunks (id, document_id, page, section, content, embedding)
                       VALUES (?, ?, ?, ?, ?, ?)""",
                    [
                        (uuid.uuid4().hex, document_id, chunk["page"], chunk["section"],
                         chunk["content"], _sparse_embedding(chunk["content"]))
                        for chunk in chunks
                    ],
                )
                connection.execute(
                    """UPDATE documents SET status='ready', page_count=?, chunk_count=?
                       WHERE id=?""",
                    (max((c["page"] or 0 for c in chunks), default=0) or None,
                     len(chunks), document_id),
                )
        except Exception as exc:
            path.unlink(missing_ok=True)
            message = str(exc) if isinstance(exc, DocumentError) else "Document processing failed."
            with self._connect() as connection:
                connection.execute(
                    "UPDATE documents SET status='failed', error=? WHERE id=?",
                    (message, document_id),
                )
            if isinstance(exc, DocumentError):
                exc.document_id = document_id
                raise
            raise DocumentError(message) from exc
        return self.document_status(owner_id, document_id)

    def delete_document(self, owner_id, document_id):
        with self._connect() as connection:
            connection.execute("PRAGMA foreign_keys=ON")
            row = connection.execute(
                "SELECT stored_name FROM documents WHERE id=? AND owner_id=?",
                (document_id, owner_id),
            ).fetchone()
            if row is None:
                return False
            connection.execute("DELETE FROM documents WHERE id=?", (document_id,))
        path = (self.upload_dir / row["stored_name"]).resolve()
        if path.parent == self.upload_dir:
            path.unlink(missing_ok=True)
        return True

    def answer(self, owner_id, question, document_ids=None, top_k=DEFAULT_TOP_K,
               min_score=DEFAULT_MIN_SCORE):
        if document_ids is None:
            documents = self.list_documents(owner_id)
            selected_ids = [doc["id"] for doc in documents if doc["status"] == "ready"]
        else:
            selected_ids = list(dict.fromkeys(document_ids))
            documents = [self.document_status(owner_id, document_id) for document_id in selected_ids]
            if any(doc is None for doc in documents):
                raise DocumentError("One or more selected documents are unavailable.")
            selected_ids = [doc["id"] for doc in documents if doc["status"] == "ready"]
        if not selected_ids:
            return {
                "answer": "The available documents do not provide sufficient information to answer this question.",
                "sources": [],
                "retrieved_count": 0,
            }

        # Summary questions: summarize the whole document instead of retrieving chunks
        if "summar" in question.lower():
            marks = ",".join("?" for _ in selected_ids)
            with self._connect() as connection:
                rows = connection.execute(
                    f"SELECT content FROM chunks WHERE document_id IN ({marks}) ORDER BY rowid",
                    selected_ids,
                ).fetchall()
            bullets = _summarize("\n".join(r["content"] for r in rows))
            return {
                "answer": "\n".join(f"• {b}" for b in bullets) or "No summary could be generated.",
                "sources": [],
                "retrieved_count": 0,
            }

        placeholders = ",".join("?" for _ in selected_ids)
        with self._connect() as connection:
            rows = connection.execute(
                f"""SELECT c.id, c.document_id, c.page, c.section, c.content, c.embedding,
                           d.filename
                    FROM chunks c JOIN documents d ON d.id=c.document_id
                    WHERE c.document_id IN ({placeholders}) AND d.owner_id=? AND d.status='ready'""",
                (*selected_ids, owner_id),
            ).fetchall()
        query_vector = _vectorizer.transform([question]).tocsr()
        query = dict(zip(query_vector.indices.tolist(), query_vector.data.tolist()))
        ranked = []
        for row in rows:
            chunk_vector = {index: value for index, value in json.loads(row["embedding"])}
            score = sum(value * chunk_vector.get(index, 0.0) for index, value in query.items())
            if score >= min_score:
                ranked.append((score, row))
        ranked.sort(key=lambda item: item[0], reverse=True)

        sources = []
        seen_content = set()
        page_counts = {}
        for score, row in ranked:
            content_key = re.sub(r"\W+", "", row["content"].lower())
            if content_key in seen_content:
                continue
            page_key = (row["document_id"], row["page"])
            if page_counts.get(page_key, 0) >= 2:
                continue
            seen_content.add(content_key)
            page_counts[page_key] = page_counts.get(page_key, 0) + 1
            sources.append({
                "id": row["id"],
                "document_id": row["document_id"],
                "filename": row["filename"],
                "page": row["page"],
                "section": row["section"],
                "score": round(float(score), 4),
                "excerpt": row["content"],
            })
            if len(sources) >= max(1, min(int(top_k), MAX_TOP_K)):
                break

        return {
            "answer": _answer_from_evidence(question, sources),
            "sources": sources,
            "retrieved_count": len(sources),
        }