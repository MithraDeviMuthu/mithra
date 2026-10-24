import io
import json
import shutil
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter

from api import app as app_module
from api.claim_extraction import ClaimDocumentStore, ClaimExtractionError
from api.rag import RagStore

try:
    import fitz  # noqa: F401
    from PIL import Image, ImageDraw, ImageFont
    OCR_READY = shutil.which("tesseract") is not None
except ImportError:
    OCR_READY = False


def _claim_image_bytes(fmt="PNG"):
    img = Image.new("RGB", (1000, 220), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=32)
    draw.text((20, 20), "Policy Number: POL-1234", fill="black", font=font)
    draw.text((20, 80), "Claim Number: CLM-5678", fill="black", font=font)
    draw.text((20, 140), "Claim Amount: $450.00", fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format=fmt)
    return buf.getvalue()


class ClaimExtractionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.orig_rag_store = app_module.RAG_STORE
        cls.orig_claim_store = app_module.CLAIM_STORE

    @classmethod
    def tearDownClass(cls):
        app_module.RAG_STORE = cls.orig_rag_store
        app_module.CLAIM_STORE = cls.orig_claim_store

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.store = ClaimDocumentStore(self.temp_dir.name)
        app_module.CLAIM_STORE = self.store
        app_module.RAG_STORE = RagStore(self.temp_dir.name)
        app_module.app.config["TESTING"] = True
        app_module.app.config["SECRET_KEY"] = "test-secret"
        self.client = app_module.app.test_client()
        with self.client.session_transaction() as sess:
            sess["user"] = "admin"          # routes now require a logged-in session

    def tearDown(self):
        try:
            self.store.close()
        except Exception:
            pass
        app_module.RAG_STORE = self.orig_rag_store
        app_module.CLAIM_STORE = self.orig_claim_store
        self.temp_dir.cleanup()

    def _pdf_bytes(self, text):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        page.merge_page(page)
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    def test_txt_extracts_claim_fields_and_uses_unavailable_for_missing(self):
        doc = self.store.upload_document(
            "owner-a",
            "claim.txt",
            b"Policy Number: POL-8841\nClaim Number: CLM-5309\nProvider Name: North Valley Clinic\nInvoice Number: INV-3342\nService Date: 2024-05-18\nClaim Amount: $1,245.50\n",
        )
        self.assertEqual(doc["status"], "ready")
        raw = json.loads(doc["raw_extraction"])
        self.assertEqual(raw["fields"]["policy_number"], "POL-8841")
        self.assertEqual(raw["fields"]["claim_number"], "CLM-5309")
        self.assertEqual(raw["fields"]["provider_name"], "North Valley Clinic")
        self.assertEqual(raw["fields"]["claim_amount"], "$1,245.50")
        self.assertIn("POL-8841", raw["source_snippets"]["policy_number"])
        self.assertIn("unavailable", raw["fields"]["discharge_date"])

    def test_txt_extracts_new_scoring_fields(self):
        doc = self.store.upload_document(
            "owner-a",
            "clinical_claim.txt",
            b"Policy Number: POL-1234\nClaim Number: CLM-9876\nPatient Age: 45\nDiagnosis Code: E11.9\nProcedure Code: 99213\nProvider Type: Clinic\nAdmission Type: Outpatient\nLength of Stay: 3\nClaim Amount: $750.00\n",
        )
        self.assertEqual(doc["status"], "ready")
        raw = json.loads(doc["raw_extraction"])["fields"]
        self.assertEqual(raw["patient_age"], "45")
        self.assertEqual(raw["diagnosis_code"], "E11.9")
        self.assertEqual(raw["procedure_code"], "99213")
        self.assertEqual(raw["provider_type"], "Clinic")
        self.assertEqual(raw["admission_type"], "Outpatient")
        self.assertEqual(raw["length_of_stay"], "3")

    def test_pdf_upload_and_review_update_preserve_raw_values(self):
        document = self.store.upload_document(
            "owner-a",
            "invoice.pdf",
            self._pdf_bytes("Policy Number POL-9001 Claim Number CLM-9912 Invoice Number INV-555 Provider Name Harbor General Hospital Claim Amount $420.00 Service Date 2024-06-11"),
        )
        raw = json.loads(document["raw_extraction"])
        self.assertEqual(raw["fields"]["claim_amount"], "$420.00")

        updated = self.store.update_document("owner-a", document["id"], {"claim_amount": "$435.00", "provider_name": "Harbor General Medical Center"})
        self.assertEqual(updated["corrected_extraction"]["fields"]["claim_amount"], "$435.00")
        self.assertEqual(updated["raw_extraction"]["fields"]["claim_amount"], "$420.00")
        self.assertEqual(updated["corrected_extraction"]["fields"]["provider_name"], "Harbor General Medical Center")

    def test_docx_extracts_fields_and_missing_fields_are_reported(self):
        document = Document()
        document.add_heading("Claim Form", level=1)
        document.add_paragraph("Policy Number: POL-4411")
        document.add_paragraph("Claim Number: CLM-6677")
        document.add_paragraph("Provider Name: Riverside Clinic")
        document.add_paragraph("Admission Date: 2024-04-02")
        document.add_paragraph("Discharge Date: 2024-04-04")
        buffer = io.BytesIO()
        document.save(buffer)

        stored = self.store.upload_document("owner-a", "claim.docx", buffer.getvalue())
        raw = json.loads(stored["raw_extraction"])
        self.assertEqual(raw["fields"]["provider_name"], "Riverside Clinic")
        self.assertEqual(raw["fields"]["claim_number"], "CLM-6677")
        self.assertIn("unavailable", raw["fields"]["claim_amount"])
        self.assertTrue(any("Missing required fields" in message for message in json.loads(stored["validation_errors"])))

    def test_invalid_empty_unsupported_and_oversized_are_rejected(self):
        with self.assertRaises(ClaimExtractionError):
            self.store.upload_document("owner-a", "bad.exe", b"not a supported file")
        with self.assertRaises(ClaimExtractionError):
            self.store.upload_document("owner-a", "empty.txt", b"")
        with self.assertRaises(ClaimExtractionError):
            self.store.upload_document("owner-a", "broken.pdf", b"not a valid pdf")
        with self.assertRaises(ClaimExtractionError):
            self.store.upload_document("owner-a", "oversized.txt", b"x" * (10 * 1024 * 1024 + 1))

    def test_invalid_image_is_rejected(self):
        with self.assertRaises(ClaimExtractionError):
            self.store.upload_document("owner-a", "scan.png", b"fake-image-content")

    @unittest.skip("OCR is disabled for image uploads.")
    def test_image_ocr_extracts_fields(self):
        doc = self.store.upload_document("owner-a", "scan.png", _claim_image_bytes("PNG"))
        self.assertEqual(doc["status"], "ready")
        fields = json.loads(doc["raw_extraction"])["fields"]
        self.assertIn("1234", fields["policy_number"])
        self.assertIn("450", fields["claim_amount"])

    @unittest.skip("OCR is disabled for scanned PDF uploads.")
    def test_scanned_pdf_falls_back_to_ocr(self):
        doc = self.store.upload_document("owner-a", "scan.pdf", _claim_image_bytes("PDF"))
        self.assertEqual(doc["status"], "ready")
        fields = json.loads(doc["raw_extraction"])["fields"]
        self.assertIn("1234", fields["policy_number"])

    def test_api_routes_upload_and_review_for_claim_documents(self):
        payload = {
            "Policy Number": "POL-55",
            "Claim Number": "CLM-55",
            "Provider Name": "Sunrise Hospital",
            "Invoice Number": "INV-55",
            "Claim Amount": "$2,000.00",
            "Service Date": "2024-01-09",
        }
        file_bytes = b"Policy Number: POL-55\nClaim Number: CLM-55\nProvider Name: Sunrise Hospital\nInvoice Number: INV-55\nClaim Amount: $2,000.00\nService Date: 2024-01-09\n"
        response = self.client.post("/api/claim-documents", data={"file": (io.BytesIO(file_bytes), "claim.txt")}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 201)
        document = response.get_json()["claim_document"]
        self.assertEqual(document["status"], "ready")

        detail = self.client.get(f"/api/claim-documents/{document['id']}")
        self.assertEqual(detail.status_code, 200)
        detail_data = detail.get_json()["claim_document"]
        self.assertEqual(detail_data["raw_extraction"]["fields"]["claim_number"], "CLM-55")

        review = self.client.post(f"/api/claim-documents/{document['id']}/review", json={"claim_amount": "$2,250.00"})
        self.assertEqual(review.status_code, 200)
        corrected = review.get_json()["claim_document"]["corrected_extraction"]["fields"]["claim_amount"]
        self.assertEqual(corrected, "$2,250.00")

    def test_api_upload_claim_document_with_rag_summary(self):
        file_bytes = b"Policy Number: POL-55\nClaim Number: CLM-55\nProvider Name: Sunrise Hospital\nInvoice Number: INV-55\nClaim Amount: $2,000.00\nService Date: 2024-01-09\nTreatment: Physical therapy follow up and consultation."
        response = self.client.post(
            "/api/claim-documents",
            data={"file": (io.BytesIO(file_bytes), "claim_rag.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201)
        doc = response.get_json()["claim_document"]
        rag_doc_id = doc.get("rag_document_id")
        self.assertIsNotNone(rag_doc_id)

        chat_res = self.client.post("/api/chat", json={
            "question": "Summarize this document",
            "document_ids": [rag_doc_id],
        })
        self.assertEqual(chat_res.status_code, 200)
        answer = chat_res.get_json().get("answer", "")
        self.assertTrue(len(answer) > 0)


if __name__ == "__main__":
    unittest.main()